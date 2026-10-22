#!/usr/bin/env python3
"""Structural QA for the revised deck.

Checks, per slide:
  * a speaker-notes block exists and its word count is in range
  * the native object mix (text boxes, autoshapes, connectors, tables, charts,
    pictures) so editability can be reported
  * every picture is an individual asset (screenshot/photo), never a full-slide
    background, and that no slide is a single flattened image
  * the smallest explicit font size used for on-slide text
"""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PATH = sys.argv[1] if len(sys.argv) > 1 else (
    "presentation/WIT_Barcode_Student_Information_System_Technical_Presentation_Revised.pptx")
NOTES_MIN, NOTES_MAX = 100, 190

prs = Presentation(PATH)
print(f"file: {PATH}")
print(f"slides: {len(prs.slides)}")
print(f"size: {prs.slide_width / 914400:.3f} x {prs.slide_height / 914400:.3f} in")

totals = {"tb": 0, "shape": 0, "conn": 0, "table": 0, "chart": 0, "pic": 0}
bad_notes = []
flat = []
min_fonts = []

for i, slide in enumerate(prs.slides, 1):
    c = {"tb": 0, "shape": 0, "conn": 0, "table": 0, "chart": 0, "pic": 0}
    pics = 0
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            c["pic"] += 1
            pics += 1
            # a picture covering ~the whole slide would be a flattened render
            if sh.width >= prs.slide_width * 0.92 and sh.height >= prs.slide_height * 0.92:
                flat.append(i)
        elif sh.has_table if hasattr(sh, "has_table") else False:
            c["table"] += 1
        elif getattr(sh, "has_chart", False):
            c["chart"] += 1
        elif sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            c["tb"] += 1
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE:
            c["conn"] += 1
        else:
            c["shape"] += 1
    for k in c:
        totals[k] += c[k]

    # smallest explicit on-slide font size
    sizes = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    sizes.append(r.font.size.pt)
    if slide.has_notes_slide:
        txt = slide.notes_slide.notes_text_frame.text.strip()
    else:
        txt = ""
    words = len(txt.split())
    if not (NOTES_MIN <= words <= NOTES_MAX):
        bad_notes.append((i, words))
    mn = min(sizes) if sizes else None
    min_fonts.append((i, mn, words))

print("\nper-slide: native objects / notes words / min on-slide font pt")
for i, (s, (_, mn, words)) in enumerate(zip(prs.slides, min_fonts), 1):
    c = {"tb": 0, "shape": 0, "conn": 0, "table": 0, "chart": 0, "pic": 0}
    for sh in s.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            c["pic"] += 1
        elif getattr(sh, "has_table", False) and sh.has_table:
            c["table"] += 1
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            c["chart"] += 1
        elif sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            c["tb"] += 1
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE:
            c["conn"] += 1
        else:
            c["shape"] += 1
    print(f"  {i:02d}  text={c['tb']:2d} shape={c['shape']:3d} conn={c['conn']:2d} "
          f"table={c['table']} chart={c['chart']} pic={c['pic']}  "
          f"notes={words:3d}w  minfont={mn}")

print("\ntotals:", totals)
print(f"slides whose notes are outside {NOTES_MIN}-{NOTES_MAX} words: {len(bad_notes)}")
for i, w in bad_notes:
    print(f"   slide {i}: {w} words")
print(f"full-slide pictures (possible flattening): {flat or 'none'}")

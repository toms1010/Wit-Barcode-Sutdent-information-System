#!/usr/bin/env python3
"""Estimate whether each editable text box can contain its text.

Uses Liberation Sans metrics (the font LibreOffice substitutes for Arial) and
the shape's real width/height. Flags boxes whose wrapped text needs more
vertical space than the shape provides.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
PTS_PER_IN = 72.0
SCALE = 4  # supersample for stable measurement


def font(size_pt: float, bold: bool):
    return ImageFont.truetype(BOLD if bold else REG, max(1, int(round(size_pt * SCALE))))


def wrap_lines(text: str, width_pt: float, size_pt: float, bold: bool):
    f = font(size_pt, bold)
    max_px = width_pt * SCALE
    lines = 0
    for para in text.split("\n"):
        words = para.split(" ")
        cur = ""
        n = 1
        for word in words:
            trial = word if not cur else cur + " " + word
            if f.getlength(trial) <= max_px or not cur:
                cur = trial
            else:
                n += 1
                cur = word
        lines += n
    return lines


def main():
    path = Path(sys.argv[1] if len(sys.argv) > 1 else
                "presentation/WIT_Barcode_Student_Information_System_Technical_Presentation.pptx")
    prs = Presentation(path)
    issues = 0
    for pno, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            tf = shape.text_frame
            w_in = Emu(shape.width).inches - Emu(tf.margin_left).inches - Emu(tf.margin_right).inches
            h_in = Emu(shape.height).inches - Emu(tf.margin_top).inches - Emu(tf.margin_bottom).inches
            if w_in <= 0.05:
                continue
            need = 0.0
            max_size = 0.0
            lines_total = 0
            for p in tf.paragraphs:
                text = "".join(r.text for r in p.runs)
                if not text:
                    continue
                size = max((r.font.size.pt for r in p.runs if r.font.size), default=12.0)
                bold = any(r.font.bold for r in p.runs)
                max_size = max(max_size, size)
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                n = wrap_lines(text, w_in * PTS_PER_IN, size, bold)
                lines_total += n
                need += n * size * 1.18 * ls / PTS_PER_IN
                need += (p.space_after.pt / PTS_PER_IN) if p.space_after else 0
            if need > h_in + 0.06:
                print(f"[slide {pno:02d}] TALL  need={need:.2f}in box={h_in:.2f}in "
                      f"size={max_size:.1f} lines={lines_total} text={tf.text[:70]!r}")
                issues += 1
    print(f"slides={len(prs.slides)} fit_issues={issues}")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
